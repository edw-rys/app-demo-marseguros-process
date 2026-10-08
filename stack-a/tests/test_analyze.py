"""Detección de tipo y extracción de texto (RF-03, RF-04)."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.extract import kind as kind_detector


# ── Magic bytes ─────────────────────────────────────────────────────────────


def test_detecta_pdf_por_magic_bytes(txt_file, tmp_path):
    """Un PDF mínimo real: la extensión no dice nada, los bytes sí."""
    pdf = tmp_path / "documento.pdf"
    pdf.write_bytes(b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n%%EOF\n")

    detected, mime = kind_detector.detect_kind(pdf)

    assert detected == "pdf"
    assert mime == "application/pdf"


def test_detecta_png_por_magic_bytes(txt_file, tmp_path):
    png = tmp_path / "escaneo.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d494844520000000100000001"
            "08060000001f15c4890000000a49444154789c6300010000050001"
            "0d0a2db40000000049454e44ae426082"
        )
    )

    detected, mime = kind_detector.detect_kind(png)

    assert detected == "png"
    assert mime == "image/png"


def test_texto_plano_cae_a_la_extension(tmp_path):
    path = tmp_path / "notas.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")

    detected, mime = kind_detector.detect_kind(path)

    assert detected == "csv"
    assert mime == "text/csv"


# ── CU-04: extensión que no corresponde al contenido ────────────────────────


def test_exe_renombrado_como_pdf_se_marca_como_mismatch(exe_disfrazado):
    detected, _ = kind_detector.detect_kind(exe_disfrazado)
    mismatch = kind_detector.has_extension_mismatch(exe_disfrazado, detected)

    assert detected == "exe"
    # Esto es lo que dispara el mensaje "no es un PDF válido" al cliente.
    assert mismatch is True


def test_extension_neutra_no_se_marca_como_mismatch(tmp_path):
    """Un `.bin` sobre un PDF no es un ataque: no debe generar alarma."""
    pdf = tmp_path / "documento.bin"
    pdf.write_bytes(b"%PDF-1.4\n%%EOF\n")

    detected, _ = kind_detector.detect_kind(pdf)

    assert kind_detector.has_extension_mismatch(pdf, detected) is False


def test_extension_correcta_no_es_mismatch(tmp_path):
    pdf = tmp_path / "factura.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%%EOF\n")

    detected, _ = kind_detector.detect_kind(pdf)

    assert kind_detector.has_extension_mismatch(pdf, detected) is False


@pytest.mark.parametrize("kind", ["pdf", "jpg", "png", "xlsx", "docx", "csv", "txt"])
def test_tipos_soportados(kind):
    assert kind_detector.is_supported(kind) is True


@pytest.mark.parametrize("kind", ["exe", "zip", "unknown"])
def test_tipos_no_soportados(kind):
    """RF-03: estos van directo a la cola de revisión."""
    assert kind_detector.is_supported(kind) is False


@pytest.mark.parametrize("kind", ["jpg", "png", "tiff", "webp"])
def test_reconoce_imagenes(kind):
    assert kind_detector.is_image(kind) is True


def test_pdf_no_es_imagen():
    """El OCR solo entra por la rama de imagen; si no, un PDF con capa de
    texto se rasterizaría y perdería Nitidez sin ganar nada."""
    assert kind_detector.is_image("pdf") is False


# ── Extracción de texto ─────────────────────────────────────────────────────


def test_extrae_texto_de_txt(txt_file):
    from app.extract import text as text_extractor

    detected, _ = kind_detector.detect_kind(txt_file)
    result = text_extractor.extract_text(txt_file, detected)

    assert "1790012345001" in result.text
    # Un .txt no pasa por OCR.
    assert result.ocr_used is False


def test_extrae_texto_de_csv(csv_file):
    from app.extract import text as text_extractor

    result = text_extractor.extract_text(csv_file, "csv")

    assert "Factura 001" in result.text
    assert result.ocr_used is False


def test_extrae_hojas_de_xlsx(tmp_path):
    from openpyxl import Workbook

    from app.extract import text as text_extractor

    path = tmp_path / "factura.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Detalle"
    ws.append(["concepto", "monto"])
    ws.append(["Factura 001", 120])
    wb.save(path)

    result = text_extractor.extract_text(path, "xlsx")

    assert "Detalle" in result.text
    assert "Factura 001" in result.text


def test_extrae_parrafos_de_docx(tmp_path):
    import docx

    from app.extract import text as text_extractor

    path = tmp_path / "solicitud.docx"
    document = docx.Document()
    document.add_paragraph("Solicitud de seguro")
    document.add_paragraph("RUC: 1790012345001")
    document.save(path)

    result = text_extractor.extract_text(path, "docx")

    assert "Solicitud de seguro" in result.text
    assert "1790012345001" in result.text


def test_respeta_max_text_chars(tmp_path, monkeypatch):
    from app.config import settings
    from app.extract import text as text_extractor

    monkeypatch.setattr(settings, "max_text_chars", 50)
    path = tmp_path / "largo.txt"
    path.write_text("palabra " * 500, encoding="utf-8")

    result = text_extractor.extract_text(path, "txt")

    assert len(result.text) <= 50